# -*- coding: utf-8 -*-
"""
tests/quality.py
────────────────
Снимок КАЧЕСТВА конвертации скан-движка (фаза 0 плана).

Прогоняет корпус (tests/corpus.py) и считает на каждый файл метрики, по которым
видно и улучшение, и регресс. LLM-доочистка ОТКЛЮЧЕНА: она недетерминирована,
а снимок должен совпадать от прогона к прогону.

Главная метрика — доля кириллических слов, известных pymorphy: она растёт, когда
распознавание становится честнее, и падает, когда текст теряется или искажается.

Использование:
    python tests/quality.py save                  # снять снимок
    python tests/quality.py check                 # сравнить с снимком
    python tests/quality.py save --groups scans   # только истинные сканы
    python tests/quality.py save --legacy-despeckle --snapshot tests/snapshots/quality-pre.json
        # воспроизвести поведение ДО п.1.1 (предохранители чистки выключены)

Допуски `check` (регресс — это падение качества, рост не ругаем):
    слов                 — не меньше 99% от снимка
    доля словарных слов  — не ниже снимка на 0.5 пп
    страниц, таблиц, ячеек — строго как в снимке
    макс. срезка чернил  — не выше снимка на 2 пп
Подсветка и мусорные блоки печатаются как информация (на них влияет п.1.4).
"""
from __future__ import annotations

import argparse
import json
import sys
import time
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # type: ignore[attr-defined]

_HERE = Path(__file__).resolve().parent
_ROOT = _HERE.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

if str(_HERE) not in sys.path:
    sys.path.insert(0, str(_HERE))

import corpus  # noqa: E402  — tests/ без __init__.py, импорт по пути

CYR = set("абвгдежзийклмнопрстуфхцчшщъыьэюяёАБВГДЕЖЗИЙКЛМНОПРСТУФХЦЧШЩЪЫЬЭЮЯЁ")
DEFAULT_SNAPSHOT = _HERE / "snapshots" / "quality.json"

# Закрытый список юр./фин. аббревиатур: нужен только для РАЗБИВКИ подсветки на
# классы (сколько жёлтого — ложная тревога). На сами метрики не влияет.
ABBR = set("""ПАО АО ООО ОАО ЗАО НАО ИП МКК МФК МФО ПКО ПКБ КПК СКПК ФНС ИФНС
УФНС МИФНС НК ГК АПК ГПК ФЗ ЕГРН ЕГРЮЛ ЕГРИП ОГРН ОГРНИП ИНН КПП СНИЛС НДФЛ НДС
ЕФРСБ РТК СРО АС ВС ВАС ФССП ГУ МВД ЗАГС ПФР СФР ЦБ РФ КНП ОКВЭД БИК""".split())


def _known_abs(row: dict) -> float:
    """Абсолютное число словарных слов из полей снимка (доля x кир-слова)."""
    return row.get("known_share", 0.0) * row.get("cyr_words", 0)


def _metrics_from_docx(path: Path) -> dict:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docling_dev.highlight import (_domain_known, _iter_paragraphs,
                                       _pymorphy_known)

    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    cells = 0
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                cells += 1
                texts.extend(p.text for p in c.paragraphs)
    words = " ".join(texts).split()
    cyr = [w.strip("«»\"'()[].,;:!?…—–-№/\\") for w in words]
    cyr = [w for w in cyr if len(w) >= 4 and all(c in CYR for c in w)]
    known = sum(1 for w in cyr
                if _pymorphy_known(w.lower()) or _domain_known(w.lower()))

    hl = [r.text.strip() for para in _iter_paragraphs(doc) for r in para.runs
          if r.font.highlight_color == WD_COLOR_INDEX.YELLOW and r.text.strip()]
    hl_core = [t.strip("«»\"'()[].,;:!?…—–-№/\\") for t in hl]
    hl_abbr = sum(1 for t in hl_core if t.upper() in ABBR
                  or (t.isupper() and 2 <= len(t) <= 6))
    hl_proper = sum(1 for t in hl_core if t and t[0].isupper() and not t.isupper()
                    and not (t.upper() in ABBR))
    from docx.oxml.ns import qn
    breaks = sum(1 for para in _iter_paragraphs(doc) for r in para.runs
                 for br in r._element.findall(qn("w:br"))
                 if br.get(qn("w:type")) == "page")
    return {
        "paragraphs": sum(1 for t in texts if t.strip()),
        "page_breaks": breaks,
        "words": len(words),
        "cyr_words": len(cyr),
        "known_share": round(known / len(cyr), 4) if cyr else 0.0,
        "tables": len(doc.tables),
        "table_cells": cells,
        "highlighted": len(hl),
        "highlighted_abbr": hl_abbr,
        "highlighted_proper": hl_proper,
        "highlighted_other": len(hl) - hl_abbr - hl_proper,
    }


def _ink_cuts(pdf: Path) -> dict:
    """Срезка чернил чисткой по страницам — метрика предохранителей п.1.1."""
    import cv2
    import fitz
    import numpy as np
    from docling_dev.ocr_preprocess import remove_stains

    cuts: list[float] = []
    with fitz.open(pdf) as d:
        for page in d:
            pix = page.get_pixmap(dpi=216)
            buf = np.frombuffer(pix.samples, dtype=np.uint8).reshape(
                pix.height, pix.width, pix.n)
            gray = cv2.cvtColor(buf, cv2.COLOR_RGB2GRAY) if pix.n >= 3 else buf[:, :, 0]
            ink0 = int((gray < 160).sum())
            cleaned = remove_stains(gray)
            ink1 = int((cleaned < 160).sum())
            cuts.append((ink0 - ink1) / max(ink0, 1) * 100)
    return {"ink_cut_max": round(max(cuts), 1) if cuts else 0.0,
            "ink_cut_mean": round(sum(cuts) / len(cuts), 1) if cuts else 0.0,
            "ink_cut_pages": [round(c, 1) for c in cuts]}


def collect(groups: str, legacy_despeckle: bool) -> dict:
    from docling_dev import build_converter, convert_pdf

    if legacy_despeckle:
        import docling_dev.ocr_preprocess as opre
        opre.MAX_INK_CUT = 1.0            # правило «вне строк» никогда не глушим
        opre._RESCUE_MIN_LETTERS = 10 ** 9  # зональное спасение никогда не срабатывает
        print("  ВНИМАНИЕ: предохранители чистки отключены (режим «до п.1.1»)")

    files = corpus.prepare(corpus.select(groups), "quality")
    if not files:
        raise SystemExit("корпус пуст: проверьте CONVERTER_CORPUS_DIR")

    conv = build_converter(ocr_preprocess=False)
    out: dict[str, dict] = {}
    for cf, pdf in files:
        docx = pdf.with_suffix(".docx")
        t0 = time.perf_counter()
        ok = convert_pdf(pdf, docx, conv, ocr_reader=None, use_word_order=False,
                         highlight=True, llm=False, ink_bold=False)
        dt = time.perf_counter() - t0
        if not ok or not docx.is_file():
            out[cf.stem] = {"group": cf.group, "error": "конвертация не удалась"}
            print(f"  ✗ {cf.stem[:46]}")
            continue
        m = {"group": cf.group, "md5": corpus.md5(pdf), "sec": round(dt, 1)}
        m.update(_metrics_from_docx(docx))
        m.update(_ink_cuts(pdf))
        out[cf.stem] = m
        print(f"  {cf.stem[:44]:44} слов={m['words']:5} словарных="
              f"{m['known_share'] * 100:5.1f}%  табл={m['tables']:2}/{m['table_cells']:3}"
              f"  подсветка={m['highlighted']:3}  срезка≤{m['ink_cut_max']:5.1f}%"
              f"  {m['sec']:5.1f}s")
    return out


def cmd_save(args) -> int:
    data = collect(args.groups, args.legacy_despeckle)
    snap = Path(args.snapshot)
    snap.parent.mkdir(parents=True, exist_ok=True)
    snap.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"\nснимок записан: {snap}  ({len(data)} файлов)")
    return 0


def cmd_check(args) -> int:
    snap = Path(args.snapshot)
    if not snap.is_file():
        raise SystemExit(f"снимка нет: {snap}. Снимите его: "
                         f"python tests/quality.py save")
    было = json.loads(snap.read_text(encoding="utf-8"))
    стало = collect(args.groups, args.legacy_despeckle)

    fails: list[str] = []
    print(f"\n{'файл':44}{'слов':>14}{'словарных':>16}{'табл/ячеек':>14}"
          f"{'подсветка':>12}{'срезка':>10}")
    for stem, now in стало.items():
        old = было.get(stem)
        if old is None:
            print(f"  {stem[:42]:42} нет в снимке — пропуск")
            continue
        if "error" in now or "error" in old:
            fails.append(f"{stem}: ошибка конвертации")
            continue

        def mark(ok: bool) -> str:
            return "" if ok else " ✗"

        ok_words = now["words"] >= old["words"] * 0.99
        # Доля словарных слов — метрика качества, но она ПАДАЕТ, когда движок
        # возвращает ранее потерянный текст: в нём есть искажения OCR (кейс п.1.1:
        # «Умерший ким клим» 830 → 993 слова и доля 93.9% → 92.8%). Поэтому
        # допуск по доле прощается, если выросло АБСОЛЮТНОЕ число словарных слов
        # при выросшем общем числе слов — это не регресс, а больше контента.
        ok_known = (now["known_share"] >= old["known_share"] - 0.005
                    or (_known_abs(now) >= _known_abs(old)
                        and now["words"] >= old["words"]))
        # Разрывы страниц — структурная метрика: их изменение означает, что
        # документ разложился по листам иначе (пустые листы, съехавшая вёрстка).
        ok_breaks = now["page_breaks"] == old.get("page_breaks", now["page_breaks"])
        ok_tables = (now["tables"] == old["tables"]
                     and now["table_cells"] == old["table_cells"])
        ok_ink = now["ink_cut_max"] <= old["ink_cut_max"] + 2.0
        print(f"  {stem[:42]:42}"
              f"{old['words']:6}→{now['words']:6}{mark(ok_words):2}"
              f"{old['known_share'] * 100:7.1f}→{now['known_share'] * 100:5.1f}%{mark(ok_known):2}"
              f"{old['tables']:4}/{old['table_cells']:<4}→{now['tables']}/{now['table_cells']}{mark(ok_tables):2}"
              f"{old['highlighted']:5}→{now['highlighted']:4}"
              f"{old['ink_cut_max']:6.1f}→{now['ink_cut_max']:5.1f}{mark(ok_ink and ok_breaks)}")
        if not ok_words:
            fails.append(f"{stem}: слов {old['words']} → {now['words']} (падение >1%)")
        if not ok_known:
            fails.append(f"{stem}: доля словарных {old['known_share'] * 100:.1f}% → "
                         f"{now['known_share'] * 100:.1f}% (падение >0.5 пп)")
        if not ok_breaks:
            fails.append(f"{stem}: разрывов страниц {old.get('page_breaks')} → "
                         f"{now['page_breaks']} (структура документа изменилась)")
        if not ok_tables:
            fails.append(f"{stem}: таблицы {old['tables']}/{old['table_cells']} → "
                         f"{now['tables']}/{now['table_cells']}")
        if not ok_ink:
            fails.append(f"{stem}: срезка чернил {old['ink_cut_max']:.1f}% → "
                         f"{now['ink_cut_max']:.1f}% (рост >2 пп)")

    missing = set(было) - set(стало)
    if missing:
        print(f"\n  в снимке есть, но не прогонялось: {', '.join(sorted(missing))}")
    if fails:
        print("\nРЕГРЕСС:")
        for f in fails:
            print(f"  ✗ {f}")
        return 1
    print("\nкачество в допусках: регресса нет")
    return 0


def main() -> int:
    p = argparse.ArgumentParser(description="Снимок качества конвертации (фаза 0)")
    p.add_argument("action", choices=("save", "check"))
    p.add_argument("--groups", default=",".join(corpus.GROUPS),
                   help="группы корпуса: scans,dirty,natives (по умолч. все)")
    p.add_argument("--snapshot", default=str(DEFAULT_SNAPSHOT))
    p.add_argument("--legacy-despeckle", action="store_true",
                   help="выключить предохранители чистки (поведение до п.1.1 плана)")
    args = p.parse_args()
    return cmd_save(args) if args.action == "save" else cmd_check(args)


if __name__ == "__main__":
    raise SystemExit(main())
